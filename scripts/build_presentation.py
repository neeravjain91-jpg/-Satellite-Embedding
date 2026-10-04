"""
Generate a professional, examiner-ready 16:9 PowerPoint presentation for the project:
'Satellite Embedding-Based Deep Learning Framework for Reconstruction of Depth-Wise Subsurface Ocean Temperature'
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="submission/Final_Project_Presentation.pptx"):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette: Deep Navy Academic Theme
    BG_DARK = RGBColor(11, 25, 44)          # #0B192C
    BG_CARD = RGBColor(19, 42, 69)          # #132A45
    TEXT_LIGHT = RGBColor(248, 250, 252)    # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)    # #94A3B8
    TEXT_DARK = RGBColor(15, 23, 42)        # #0F172A
    ACCENT_CYAN = RGBColor(56, 189, 248)    # #38BDF8
    ACCENT_BLUE = RGBColor(14, 165, 233)    # #0EA5E9
    ACCENT_GREEN = RGBColor(16, 185, 129)   # #10B981
    ACCENT_AMBER = RGBColor(245, 158, 11)   # #F59E0B
    BORDER_COLOR = RGBColor(30, 58, 95)     # #1E3A5F

    blank_layout = prs.slide_layouts[6] # blank layout

    def add_bg(slide, dark=True):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = BG_DARK if dark else RGBColor(255, 255, 255)
        shape.line.fill.background()
        return shape

    def add_header(slide, title, category="RESEARCH SUBMISSION", dark=True):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN if dark else ACCENT_BLUE

        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
        tf = t_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT if dark else TEXT_DARK

    def add_card(slide, left, top, width, height, bg_color=BG_CARD, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        return shape

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, dark=True)

    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.15), Inches(4.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_CYAN
    bar.line.fill.background()

    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11.2), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "OCEAN DEEP LEARNING & SATELLITE ALTIMETRY"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN

    p1 = tf.add_paragraph()
    p1.text = "Satellite Embedding-Based Deep Learning Framework for Reconstruction of Depth-Wise Subsurface Ocean Temperature"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT

    p2 = tf.add_paragraph()
    p2.text = "End-to-End Vertical Thermal Structure Inversion across the North Indian Ocean (0–1000 m)"
    p2.font.size = Pt(16)
    p2.font.color.rgb = TEXT_MUTED

    card_meta = add_card(s1, Inches(1.2), Inches(5.0), Inches(11.0), Inches(1.5), bg_color=BG_CARD)
    tb_m = s1.shapes.add_textbox(Inches(1.4), Inches(5.1), Inches(10.6), Inches(1.3))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True

    pm = tf_m.paragraphs[0]
    pm.text = "KEY METRICS & EXPERIMENTAL PROTOCOL"
    pm.font.size = Pt(11)
    pm.font.bold = True
    pm.font.color.rgb = ACCENT_GREEN

    pm1 = tf_m.add_paragraph()
    pm1.text = "• Domain: North Indian Ocean (5°N–30°N, 45°E–105°E) | Grid: Canonical 0.25° × 0.25° (101 × 241 cells, 15 Depths)\n" \
               "• Certified Test RMSE: 0.9800 °C (B8 Spatiotemporal Model) | Error Reduction vs. Climatology (B1): 22.11%\n" \
               "• Controls: 6-Day Purge Buffers (Zero Leakage) | Train-Only Normalization | Strict Chronological 2020 Partition (Test: Days 313–365)"
    pm1.font.size = Pt(12)
    pm1.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 2: Problem Statement & Scientific Motivation
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2, dark=True)
    add_header(s2, "Problem Statement & Scientific Motivation", "THE SCIENTIFIC CHALLENGE")

    c_w = Inches(3.7)
    c_h = Inches(5.0)
    top_pos = Inches(1.8)

    # Card 1: Ocean Opacity
    add_card(s2, Inches(0.8), top_pos, c_w, c_h)
    tb = s2.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), c_w - Inches(0.4), c_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "The Opacity Challenge"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p2 = tf.add_paragraph()
    p2.text = "\n• Seawater strongly absorbs electromagnetic radiation: satellite radiometers and scatterometers observe exclusively the skin and sub-skin (<1 mm to a few cm).\n\n" \
              "• The 3D subsurface thermal interior (>10 m to 1000 m) is invisible to direct satellite remote sensing.\n\n" \
              "• Subsurface temperature dictates ocean heat content, acoustic ducts, tropical cyclone intensification, and monsoon dynamics."
    p2.font.size = Pt(12); p2.font.color.rgb = TEXT_LIGHT

    # Card 2: In-Situ Sparsity
    add_card(s2, Inches(4.8), top_pos, c_w, c_h)
    tb = s2.shapes.add_textbox(Inches(5.0), top_pos + Inches(0.2), c_w - Inches(0.4), c_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "In-Situ Observational Sparsity"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p2 = tf.add_paragraph()
    p2.text = "\n• Autonomous Argo CTD profiling floats collect in-situ soundings only once every ~10 days per float.\n\n" \
              "• Average spatial separation between floats is ~300 km: insufficient for mesoscale eddy dynamics (<100 km).\n\n" \
              "• Research vessels and moored buoys (RAMA) offer sparse discrete points. A synoptic, continuous 3D field is physically unattainable from in-situ measurements alone."
    p2.font.size = Pt(12); p2.font.color.rgb = TEXT_LIGHT

    # Card 3: Deep Representation Solution
    add_card(s2, Inches(8.8), top_pos, c_w, c_h)
    tb = s2.shapes.add_textbox(Inches(9.0), top_pos + Inches(0.2), c_w - Inches(0.4), c_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "The Deep Learning Solution"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = ACCENT_GREEN
    p2 = tf.add_paragraph()
    p2.text = "\n• Surface variables contain hydrodynamic signatures of subsurface dynamics (SSH geostrophy, SST heating, SSS buoyancy, Ekman pumping).\n\n" \
              "• Inversion via Spatiotemporal Embedding: Learn nonlinear mappings from 5-day causal multi-satellite patches to 15 discrete vertical temperature levels.\n\n" \
              "• Provides daily, high-resolution (0.25°), synoptic 3D temperature fields across the full basin against GLORYS reanalysis."
    p2.font.size = Pt(12); p2.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 3: Study Domain & 7 Multi-Satellite Predictors
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3, dark=True)
    add_header(s3, "Study Domain & 7 Multi-Satellite Input Predictors", "DATASET SPECIFICATION")

    add_card(s3, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    tb = s3.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.1), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "North Indian Ocean Domain"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p2 = tf.add_paragraph()
    p2.text = "\n• Spatial Bounding Box: 5.00°N to 30.00°N, 45.00°E to 105.00°E\n" \
              "• Canonical Grid: 0.25° × 0.25° (101 lat × 241 lon = 24,341 horizontal cells)\n" \
              "• Valid Ocean Sea Surface Cells: 16,076 cells\n" \
              "• Active 3D Ocean-Depth Cells: 166,400 cells (GEBCO bathymetry cutoff enforced)\n" \
              "• Temporal Extent: Full Year 2020 (366 consecutive calendar days, leap day verified)\n" \
              "• 15 Target Depths: [0, 5, 10, 20, 30, 50, 75, 100, 125, 150, 200, 300, 500, 700, 1000] m\n" \
              "• Target Reference: CMEMS GLORYS12V1 (1/12° daily reanalysis, reference state)"
    p2.font.size = Pt(11); p2.font.color.rgb = TEXT_LIGHT

    add_card(s3, Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.0))
    tb = s3.shapes.add_textbox(Inches(5.8), Inches(2.0), Inches(6.5), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Harmonized Multi-Satellite Inputs"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN

    rows, cols = 8, 4
    t_shape = s3.shapes.add_table(rows, cols, Inches(5.8), Inches(2.7), Inches(6.5), Inches(3.8))
    table = t_shape.table
    table.columns[0].width = Inches(1.5)
    table.columns[1].width = Inches(1.8)
    table.columns[2].width = Inches(1.4)
    table.columns[3].width = Inches(1.8)

    headers = ["Variable", "Source Product", "Sensor Type", "Physical Role"]
    for c_idx, h in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid(); cell.fill.fore_color.rgb = BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h; p.font.bold = True; p.font.size = Pt(10); p.font.color.rgb = ACCENT_CYAN

    data = [
        ("SST", "OSTIA L4 Reprocessed", "IR + Microwave", "Upper boundary thermal constraint"),
        ("SSS", "Copernicus Multi-Obs", "SMOS / In-situ", "Salinity & halocline buoyancy"),
        ("SSH / SLA", "DUACS Multi-Mission", "Radar Altimeter", "Baroclinic mode & thermocline tilt"),
        ("Current U", "OSCAR Ocean Currents", "Altimetry + Wind", "Zonal geostrophic & Ekman advection"),
        ("Current V", "OSCAR Ocean Currents", "Altimetry + Wind", "Meridional boundary currents"),
        ("Wind U", "CCMP V3.1 10m Wind", "Radiom. + Scatterom.", "Surface zonal shear stress"),
        ("Wind V", "CCMP V3.1 10m Wind", "Radiom. + Scatterom.", "Ekman pumping & upwelling"),
    ]
    for r_idx, row in enumerate(data, start=1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.fill.solid(); cell.fill.fore_color.rgb = BG_CARD if r_idx % 2 == 0 else RGBColor(25, 52, 85)
            p = cell.text_frame.paragraphs[0]
            p.text = val; p.font.size = Pt(9); p.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 4: Leakage Controls & Experimental Protocol
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4, dark=True)
    add_header(s4, "Experimental Protocol & Scientific Leakage Controls", "RIGOROUS METHODOLOGY")

    add_card(s4, Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.8))
    tb = s4.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.3), Inches(1.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Strict Chronological Partitioning (Full-Year 2020 Calendar)"
    p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN

    p1 = tf.add_paragraph()
    p1.text = "TRAIN SET (Days 0–252: Jan 1 – Sep 9)   |   PURGE 1 (Days 253–258)   |   VAL SET (Days 259–306: Sep 16 – Nov 2)   |   PURGE 2 (Days 307–312)   |   TEST SET (Days 313–365: Nov 9 – Dec 31)\n" \
              "  253 Days (~69.1% of calendar)              6-Day Purge Buffer                48 Days (~13.1%)                         6-Day Purge Buffer                 53 Days (~14.5%)"
    p1.font.size = Pt(11); p1.font.color.rgb = TEXT_LIGHT; p1.font.bold = True

    p2 = tf.add_paragraph()
    p2.text = "✓ Temporal purge buffers (6 days) strictly exceed the 5-day causal input window: guarantees zero historical or future contamination across splits."
    p2.font.size = Pt(10.5); p2.font.color.rgb = ACCENT_GREEN

    bc_w = Inches(3.7)
    bc_h = Inches(3.1)
    bc_top = Inches(3.8)

    # 1. Zero-Leakage Normalization
    add_card(s4, Inches(0.8), bc_top, bc_w, bc_h)
    tb = s4.shapes.add_textbox(Inches(1.0), bc_top + Inches(0.2), bc_w - Inches(0.4), bc_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. Train-Only Normalization"
    p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p2 = tf.add_paragraph()
    p2.text = "\n• Mean & standard deviation fitted strictly on training days (0–252).\n" \
              "• Frozen statistics saved to JSON and applied to validation and test sets.\n" \
              "• Absolutely zero data leakage from future evaluation periods."
    p2.font.size = Pt(11); p2.font.color.rgb = TEXT_LIGHT

    # 2. Causal Receptive Field
    add_card(s4, Inches(4.8), bc_top, bc_w, bc_h)
    tb = s4.shapes.add_textbox(Inches(5.0), bc_top + Inches(0.2), bc_w - Inches(0.4), bc_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. Causal Slicing"
    p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p2 = tf.add_paragraph()
    p2.text = "\n• Predictions at day t use only past days: [t-4, t-3, t-2, t-1, t].\n" \
              "• No forward temporal padding or centered filters.\n" \
              "• Strictly mimics operational real-time forecasting conditions."
    p2.font.size = Pt(11); p2.font.color.rgb = TEXT_LIGHT

    # 3. Canonical Ocean Masking
    add_card(s4, Inches(8.8), bc_top, bc_w, bc_h)
    tb = s4.shapes.add_textbox(Inches(9.0), bc_top + Inches(0.2), bc_w - Inches(0.4), bc_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "3. Bathymetric Cutoffs"
    p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p2 = tf.add_paragraph()
    p2.text = "\n• Sub-seafloor cells strictly masked as NaN via GEBCO bathymetry.\n" \
              "• Loss computed only on active ocean depths.\n" \
              "• Prevents models from hallucinating temperature inside solid continental crust."
    p2.font.size = Pt(11); p2.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 5: Baseline Hierarchy Progression (B0 to B8)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5, dark=True)
    add_header(s5, "Systematic Baseline Hierarchy Progression (B0 – B8)", "EXPERIMENTAL TAXONOMY")

    add_card(s5, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tb = s5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Four-Tier Progression from Physical Baselines to Deep Spatiotemporal Networks"
    p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN

    rows, cols = 5, 5
    t_shape = s5.shapes.add_table(rows, cols, Inches(1.0), Inches(2.6), Inches(11.3), Inches(3.9))
    table = t_shape.table
    table.columns[0].width = Inches(1.2)
    table.columns[1].width = Inches(2.5)
    table.columns[2].width = Inches(2.6)
    table.columns[3].width = Inches(2.6)
    table.columns[4].width = Inches(2.4)

    tier_headers = ["Tier", "Models", "Input Context", "Key Scientific Hypothesis", "Observed Outcome"]
    for c_idx, h in enumerate(tier_headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid(); cell.fill.fore_color.rgb = BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h; p.font.bold = True; p.font.size = Pt(10.5); p.font.color.rgb = ACCENT_CYAN

    tier_data = [
        ("Level 0\nPhysical", "B0: Day-0 Persistence\nB0b: Day-252 Persistence\nB1: Daily Climatology", "Target column history\n(No satellite features)", "Tests whether ocean state is trivial or governed by static climatological cycles.", "Climatology (1.2582 °C) outperforms persistence (1.5220 °C) over long evaluation horizons."),
        ("Level 1\nTabular ML", "B2: Ridge (alpha=100,000)\nB3: Random Forest\nB4: LightGBM", "Pointwise surface 7-features (1x1 pixel, 1 day)", "Tests whether linear/nonlinear surface-to-subsurface coupling beats climatology.", "All tabular ML models beat B1 (~1.028–1.045 °C). Surface coupling provides strong predictive signals."),
        ("Level 2\nDL Ablations", "B5: Pointwise MLP\nB6: Spatial CNN (3x3)\nB7: Temporal GRU (5-day)", "Isolated spatial OR temporal context", "Tests isolated spatial vs isolated temporal context in neural architectures.", "Uncoupled DL suffers from overfitting (B5: 1.5524 °C, B7: 1.5320 °C). Spatial CNN (1.2702 °C) helps but lacks time."),
        ("Level 3\nFull DL", "B8: Spatiotemporal Embedding Network", "5-day causal window × 3x3 spatial patch × 7 vars", "Tests joint spatiotemporal embedding and latent space inversion.", "B8 achieves 0.9800 °C (22.11% gain vs B1, 4.74% gain vs B4). Spatiotemporal synergy is decisive.")
    ]
    for r_idx, row in enumerate(tier_data, start=1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.fill.solid(); cell.fill.fore_color.rgb = BG_CARD if r_idx % 2 == 0 else RGBColor(25, 52, 85)
            p = cell.text_frame.paragraphs[0]
            p.text = val; p.font.size = Pt(9.5); p.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 6: B8 Architecture & 128-D Latent Space
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6, dark=True)
    add_header(s6, "B8 Spatiotemporal Architecture & 128-D Latent Space", "DEEP LEARNING MODEL DESIGN")

    sc_w = Inches(3.7)
    sc_h = Inches(5.0)
    sc_top = Inches(1.8)

    # 1. Spatial CNN
    add_card(s6, Inches(0.8), sc_top, sc_w, sc_h)
    tb = s6.shapes.add_textbox(Inches(1.0), sc_top + Inches(0.2), sc_w - Inches(0.4), sc_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Stage 1: Spatial CNN"
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p2 = tf.add_paragraph()
    p2.text = "\n• Input: Time-distributed 3 × 3 spatial patches with 7 surface channels.\n\n" \
              "• Layer 1: Conv2D(7 → 32, kernel=3, pad=1)\n" \
              "  + BatchNorm2D + ReLU\n\n" \
              "• Layer 2: Conv2D(32 → 64, kernel=3, pad=1)\n" \
              "  + BatchNorm2D + ReLU\n\n" \
              "• AdaptiveAvgPool2D((1, 1)) + Flatten\n\n" \
              "• Output: 64-dimensional spatial feature vector s_t for each day t ∈ {1..5}."
    p2.font.size = Pt(11); p2.font.color.rgb = TEXT_LIGHT

    # 2. Temporal GRU
    add_card(s6, Inches(4.8), sc_top, sc_w, sc_h)
    tb = s6.shapes.add_textbox(Inches(5.0), sc_top + Inches(0.2), sc_w - Inches(0.4), sc_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Stage 2: 2-Layer GRU"
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p2 = tf.add_paragraph()
    p2.text = "\n• Input: Sequence of spatial vectors [s_1, s_2, s_3, s_4, s_5] ∈ ℝ^{5 × 64}.\n\n" \
              "• Recurrent Core: 2-layer Gated Recurrent Unit (input_size=64, hidden_size=128, batch_first=True).\n\n" \
              "• Latent Bottleneck: LayerNorm(128) applied to terminal hidden state h_5 ∈ ℝ^{128}.\n\n" \
              "• Produces 128-dimensional compressed ocean latent embedding."
    p2.font.size = Pt(11); p2.font.color.rgb = TEXT_LIGHT

    # 3. Vertical MLP Decoder
    add_card(s6, Inches(8.8), sc_top, sc_w, sc_h)
    tb = s6.shapes.add_textbox(Inches(9.0), sc_top + Inches(0.2), sc_w - Inches(0.4), sc_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Stage 3: Depth Decoder"
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = ACCENT_GREEN
    p2 = tf.add_paragraph()
    p2.text = "\n• Input: 128-D Ocean Latent Vector.\n\n" \
              "• Projection Layer 1: Linear(128 → 64) + ReLU\n\n" \
              "• Projection Layer 2: Linear(64 → 15)\n\n" \
              "• Output: Predicted temperature profile across all 15 canonical depths.\n\n" \
              "• Total Model Parameters: 203,791 (0.81 MB memory footprint).\n\n" \
              "• Lightweight architecture suitable for efficient inference."
    p2.font.size = Pt(11); p2.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 7: Master Benchmark Results
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7, dark=True)
    add_header(s7, "Master Benchmark Results & Statistical Significance", "CERTIFIED EVALUATION METRICS")

    add_card(s7, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.3), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Independent Test Partition Evaluation (Days 313–365, Full Basin)"
    p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN

    rows, cols = 11, 6
    t_shape = s7.shapes.add_table(rows, cols, Inches(1.0), Inches(2.4), Inches(11.3), Inches(4.2))
    table = t_shape.table
    table.columns[0].width = Inches(1.1)
    table.columns[1].width = Inches(3.2)
    table.columns[2].width = Inches(1.6)
    table.columns[3].width = Inches(1.8)
    table.columns[4].width = Inches(1.8)
    table.columns[5].width = Inches(1.8)

    bm_headers = ["Model ID", "Model Family / Description", "Parameters", "Test RMSE (°C)", "Gain vs B1 (%)", "Protocol Status"]
    for c_idx, h in enumerate(bm_headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid(); cell.fill.fore_color.rgb = BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h; p.font.bold = True; p.font.size = Pt(10); p.font.color.rgb = ACCENT_CYAN

    bm_data = [
        ("B0", "Day-0 Persistence", "0", "1.5220", "-20.97%", "Locked Baseline"),
        ("B0b", "Day-252 Persistence", "0", "1.7287", "-37.40%", "Locked Baseline"),
        ("B1", "Daily Mean Climatology", "0", "1.2582", "Reference", "Locked Reference"),
        ("B2", "Ridge Regression (alpha=100,000)", "120", "1.0295", "+18.18%", "Locked Baseline"),
        ("B3", "Random Forest Regressor (100 trees)", "~850,000", "1.0452", "+16.93%", "Locked Baseline"),
        ("B4", "LightGBM Gradient Boosting", "~320,000", "1.0288", "+18.23%", "Locked Baseline"),
        ("B5", "Pointwise MLP (3 Hidden Layers)", "26,767", "1.5524", "-23.38%", "Locked Baseline"),
        ("B6", "Spatial CNN (3x3 Patches)", "30,991", "1.2702", "-0.95%", "Locked Baseline"),
        ("B7", "Temporal GRU (5-Day Causal)", "44,111", "1.5320", "-21.76%", "Locked Baseline"),
        ("B8", "Spatiotemporal Embedding Network", "203,791", "0.9800", "+22.11%", "BEST INTERNAL")
    ]
    for r_idx, row in enumerate(bm_data, start=1):
        is_b8 = (row[0] == "B8")
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.fill.solid()
            if is_b8:
                cell.fill.fore_color.rgb = RGBColor(16, 68, 60)
            else:
                cell.fill.fore_color.rgb = BG_CARD if r_idx % 2 == 0 else RGBColor(25, 52, 85)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(9.5)
            p.font.bold = is_b8
            p.font.color.rgb = ACCENT_GREEN if (is_b8 and c_idx >= 3) else (TEXT_LIGHT)

    # ==========================================
    # SLIDE 8: Depth-Wise Stratification & Thermocline Decomposition
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8, dark=True)
    add_header(s8, "Depth-Wise Stratification & Thermocline Error Peak", "PHYSICAL OCEANOGRAPHIC ANALYSIS")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tb = s8.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Certified B8 Depth-Wise Test RMSE"
    p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN

    rows, cols = 16, 3
    t_shape = s8.shapes.add_table(rows, cols, Inches(1.0), Inches(2.6), Inches(5.2), Inches(4.0))
    table = t_shape.table
    table.columns[0].width = Inches(1.4)
    table.columns[1].width = Inches(1.6)
    table.columns[2].width = Inches(2.2)

    headers = ["Depth (m)", "B8 RMSE (°C)", "Oceanographic Regime"]
    for c_idx, h in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid(); cell.fill.fore_color.rgb = BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h; p.font.bold = True; p.font.size = Pt(9.5); p.font.color.rgb = ACCENT_CYAN

    depth_data = [
        ("0 m", "0.4369", "Surface Boundary Layer"),
        ("5 m", "0.4381", "Epipelagic Mixed Layer"),
        ("10 m", "0.4635", "Epipelagic Mixed Layer"),
        ("20 m", "0.5962", "Mixed Layer Base"),
        ("30 m", "0.8607", "Upper Thermocline Transition"),
        ("50 m", "1.3493", "Main Thermocline Gradient"),
        ("75 m", "1.8110", "★ Peak Error (Max Stratification)"),
        ("100 m", "1.7651", "Sub-Thermocline Shear"),
        ("125 m", "1.5022", "Permanent Pycnocline"),
        ("150 m", "1.3650", "Permanent Pycnocline"),
        ("200 m", "1.1834", "Mesopelagic Transition"),
        ("300 m", "0.9845", "Mesopelagic Zone"),
        ("500 m", "0.6870", "Deep Quiescent Water Mass"),
        ("700 m", "0.6619", "Deep Ocean"),
        ("1000 m", "0.5959", "Deep Stable Ocean"),
    ]
    for r_idx, (d, r, reg) in enumerate(depth_data, start=1):
        is_peak = ("★" in reg)
        for c_idx, val in enumerate([d, r, reg]):
            cell = table.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(70, 20, 20) if is_peak else (BG_CARD if r_idx % 2 == 0 else RGBColor(25, 52, 85))
            p = cell.text_frame.paragraphs[0]
            p.text = val; p.font.size = Pt(8.5); p.font.bold = is_peak
            p.font.color.rgb = ACCENT_AMBER if is_peak else TEXT_LIGHT

    add_card(s8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    tb = s8.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Key Scientific Insights"
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN

    p2 = tf.add_paragraph()
    p2.text = "\n1. Surface Anchoring (0–10 m: RMSE ~0.44 °C):\n" \
              "   Satellite SST directly constrains epipelagic temperatures, yielding superior fidelity near the boundary.\n\n" \
              "2. Thermocline Peak Variance (50–125 m: Peak 1.8110 °C at 75 m):\n" \
              "   The thermocline exhibits steep vertical temperature gradients (dT/dz > 0.1 °C/m) and energetic internal wave oscillations. Small vertical displacements translate to large thermal discrepancies.\n\n" \
              "3. Deep Quiescent Stability (500–1000 m: RMSE ~0.60 °C):\n" \
              "   Mesopelagic waters have low ambient thermal variance (σ < 0.8 °C). The network captures this stable baseline.\n\n" \
              "4. Critical Reporting Requirement:\n" \
              "   The overall 0.9800 °C is a column average. Stating 0.9800 °C uniformly across depths is scientifically invalid."
    p2.font.size = Pt(11); p2.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 9: Regional & Seasonal Basin Performance
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9, dark=True)
    add_header(s9, "Regional Basin & Seasonal Performance Breakdown", "HYDRODYNAMIC HETEROGENEITY")

    add_card(s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tb = s9.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Regional Cosine-Latitude Weighted RMSE"
    p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN

    p2 = tf.add_paragraph()
    p2.text = "\n• Full Domain (NIO): 0.9642 °C\n\n" \
              "• Arabian Sea (AS): 1.0907 °C\n" \
              "  - Impacted by intense evaporative cooling, winter convective overturn, and vigorous eddy kinetic energy from the western boundary Somali current system.\n\n" \
              "• Bay of Bengal (BoB): 0.6775 °C\n" \
              "  - Significantly lower error due to massive Ganges-Brahmaputra river freshwater influx creating a strongly stratified 'barrier layer' that stabilizes upper-layer thermal profiles."
    p2.font.size = Pt(11.5); p2.font.color.rgb = TEXT_LIGHT

    add_card(s9, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    tb = s9.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Certified Seasonal Test Subsets"
    p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN

    p2 = tf.add_paragraph()
    p2.text = "\n• Late Fall (Nov 09 – Nov 30): 1.0059 °C\n" \
              "  - Post-monsoon transition characterized by retreating southwest monsoon currents and decaying anticyclonic eddies.\n\n" \
              "• Early Winter (Dec 01 – Dec 31): 0.9060 °C\n" \
              "  - Established Northeast Monsoon circulation with steady northeasterly wind stress and stable thermal stratification.\n\n" \
              "• Reporting Boundary Disclosure:\n" \
              "  - SW Monsoon and Pre-Monsoon occur entirely within the Training split (Days 0–252).\n" \
              "  - In adherence to protocol, only Late Fall and Early Winter exist in the certified Test partition (Days 313–365)."
    p2.font.size = Pt(11.5); p2.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 10: Interactive UI Prototype
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10, dark=True)
    add_header(s10, "Interactive Web Prototype & System Architecture", "DEPLOYMENT & SCIENTIFIC VISUALIZATION")

    ic_w = Inches(3.7)
    ic_h = Inches(5.0)
    ic_top = Inches(1.8)

    # 1. Tech Stack
    add_card(s10, Inches(0.8), ic_top, ic_w, ic_h)
    tb = s10.shapes.add_textbox(Inches(1.0), ic_top + Inches(0.2), ic_w - Inches(0.4), ic_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Frontend Architecture"
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p2 = tf.add_paragraph()
    p2.text = "\n• Framework: React 18 with TypeScript and Vite bundling.\n\n" \
              "• Styling: Tailwind CSS responsive utility layout with dark oceanic styling.\n\n" \
              "• Charts: Interactive Recharts SVG rendering for depth profiles, soundings, and benchmark comparisons.\n\n" \
              "• Decoupled Architecture: 100% client-side simulation running on local mock data (zero external API dependencies)."
    p2.font.size = Pt(11); p2.font.color.rgb = TEXT_LIGHT

    # 2. Key Interactive Capabilities
    add_card(s10, Inches(4.8), ic_top, ic_w, ic_h)
    tb = s10.shapes.add_textbox(Inches(5.0), ic_top + Inches(0.2), ic_w - Inches(0.4), ic_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Interactive Features"
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p2 = tf.add_paragraph()
    p2.text = "\n• Depth Sounding Slider: Real-time dynamic slicing across all 15 discrete ocean depths (0–1000 m).\n\n" \
              "• Regional Basin Selector: Dynamic inspection of Full Domain, Arabian Sea, and Bay of Bengal metrics.\n\n" \
              "• Model Comparison: Side-by-side reconstruction error inspection across B0–B8.\n\n" \
              "• Profile Explorer: Simulated CTD vertical sounding curves comparing B8 predictions against reference states."
    p2.font.size = Pt(11); p2.font.color.rgb = TEXT_LIGHT

    # 3. Live Production Deployment
    add_card(s10, Inches(8.8), ic_top, ic_w, ic_h)
    tb = s10.shapes.add_textbox(Inches(9.0), ic_top + Inches(0.2), ic_w - Inches(0.4), ic_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Live Deployment"
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = ACCENT_GREEN
    p2 = tf.add_paragraph()
    p2.text = "\n• Production URL:\n" \
              "  https://code-gules-three.vercel.app\n\n" \
              "• Global CDN: Deployed on Vercel with zero latency and edge caching.\n\n" \
              "• Quality Assurance:\n" \
              "  - 0 TypeScript compilation errors.\n" \
              "  - Rigorous scientific labels: 'B0b Day-252 Persistence', 'PROTOTYPE • LOCAL SIMULATION'.\n" \
              "  - Complete adherence to certified benchmark values."
    p2.font.size = Pt(11); p2.font.color.rgb = TEXT_LIGHT

    # ==========================================
    # SLIDE 11: Limitations, Future Work & Conclusion
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_bg(s11, dark=True)
    add_header(s11, "Scientific Disclosures, Future Directions & Conclusions", "SUMMARY & IMPACT")

    cc_w = Inches(3.7)
    cc_h = Inches(5.0)
    cc_top = Inches(1.8)

    # 1. Limitations
    add_card(s11, Inches(0.8), cc_top, cc_w, cc_h)
    tb = s11.shapes.add_textbox(Inches(1.0), cc_top + Inches(0.2), cc_w - Inches(0.4), cc_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Scientific Limitations"
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = ACCENT_AMBER
    p2 = tf.add_paragraph()
    p2.text = "\n• GLORYS Reanalysis Reference: Target is a numerical assimilation product, not direct in-situ ground truth.\n\n" \
              "• Deep Ocean Variance: Low RMSE (<0.7 °C) at 500–1000m reflects low ambient physical variance, not deeper penetration.\n\n" \
              "• Single-Year Scope (2020): Models trained on 2020 may not fully generalize across decadal climate modes (IOD, ENSO).\n\n" \
              "• Argo Assimilation: Standard Argo profiles are assimilated into GLORYS."
    p2.font.size = Pt(10.5); p2.font.color.rgb = TEXT_LIGHT

    # 2. Future Work
    add_card(s11, Inches(4.8), cc_top, cc_w, cc_h)
    tb = s11.shapes.add_textbox(Inches(5.0), cc_top + Inches(0.2), cc_w - Inches(0.4), cc_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Future Directions"
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p2 = tf.add_paragraph()
    p2.text = "\n• Decadal Pretraining: Scaling to multi-decade training (1993–2020) to capture interannual modes.\n\n" \
              "• Withheld Float Validation: Evaluating against independently withheld Argo floats or dedicated research cruise CTDs.\n\n" \
              "• Physical Inductive Biases: Incorporating hydrostatic balance, geostrophic constraints, and buoyancy loss functions into the neural architecture.\n\n" \
              "• Foundation Ocean Model: Pretraining masked autoencoders on 3D reanalyses."
    p2.font.size = Pt(10.5); p2.font.color.rgb = TEXT_LIGHT

    # 3. Final Conclusion
    add_card(s11, Inches(8.8), cc_top, cc_w, cc_h)
    tb = s11.shapes.add_textbox(Inches(9.0), cc_top + Inches(0.2), cc_w - Inches(0.4), cc_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Conclusion"
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = ACCENT_GREEN
    p2 = tf.add_paragraph()
    p2.text = "\n• B8 Spatiotemporal Embedding achieves 0.9800 °C overall column-averaged test RMSE (22.11% gain vs climatology).\n\n" \
              "• Spatial context + causal temporal memory captures predictive relationships: uncoupled architectures fail to beat tabular baselines.\n\n" \
              "• Complete scientific reproducibility achieved: 97/97 tests passing, zero data leakage, and interactive live UI prototype.\n\n" \
              "• Ready for submission and academic defence."
    p2.font.size = Pt(10.5); p2.font.color.rgb = TEXT_LIGHT

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path} ({len(prs.slides)} slides)")

if __name__ == "__main__":
    create_deck()
